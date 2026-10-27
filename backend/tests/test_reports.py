from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app.engine.costing import CostingRequest, build_bom
from app.financial.dcf_model import FinancialInputs, calculate_dcf, monte_carlo
from app.main import app, assessment, demo_assessment
from app.reports.excel_builder import build_excel
from app.reports.pdf_builder import build_pdf

client = TestClient(app)


@pytest.fixture
def result():
    return assessment(demo_assessment())


def test_workbook_financial_cells_are_formulas(result):
    wb = load_workbook(BytesIO(build_excel(result.bom, result.financial)))
    assert wb.sheetnames == ["Summary", "Inputs", "Equipment BOM", "Cash Flow", "Notes"]
    assert wb.calculation.fullCalcOnLoad
    assert not wb.security.lockStructure
    assert wb["Cash Flow"].max_row == 106
    for row in wb["Cash Flow"].iter_rows(min_row=6, max_row=106, max_col=19):
        assert all(cell.data_type == "f" for cell in row)
    assert (
        wb["Summary"]["C6"].value
        == "=NPV(Inputs!C7,'Cash Flow'!L7:L106)+'Cash Flow'!L6"
    )
    assert "IRR(" in wb["Summary"]["C7"].value
    assert "OFFSET(" in wb["Summary"]["C7"].value
    assert "Inputs!C5=1" in wb["Summary"]["C7"].value
    assert wb["Inputs"]["C17"].value == "='Equipment BOM'!T10"
    assert (
        wb["Cash Flow"]["G7"].value
        == "=IF($A7<=Inputs!$C$5,IF($A7<=Inputs!$C$6,Inputs!$C$17/Inputs!$C$6,0),0)"
    )
    assert wb["Equipment BOM"]["T10"].value == "=SUM(T6:T9)"
    assert not wb._external_links


def test_pdf_is_two_pages_with_key_results(result):
    risk = result.uncertainty or monte_carlo(result.financial.inputs)
    content = build_pdf(result.bom, result.financial, risk)
    reader = PdfReader(BytesIO(content))
    assert len(reader.pages) == 2
    text = "\n".join(page.extract_text() for page in reader.pages)
    for token in (
        "Executive summary",
        "IRR",
        "LCOP",
        "P10",
        "P50",
        "P90",
        "Risk",
        "CEPCI",
    ):
        assert token in text
    assert f"{result.financial.npv_usd:,.0f}" in text


def test_large_bom_still_two_pages_and_literal_excel_names():
    payload = demo_assessment().costing.model_dump()
    base = payload["equipment"][0]
    payload["equipment"] = [base | {"unit_id": f"=2+{i}"} for i in range(150)]
    bom = build_bom(CostingRequest.model_validate(payload))
    financial = calculate_dcf(
        FinancialInputs(
            total_capex_usd=bom.total_bare_module_cost_usd,
            **demo_assessment().financial.model_dump(),
        )
    )
    pdf = build_pdf(bom, financial, monte_carlo(financial.inputs))
    assert len(PdfReader(BytesIO(pdf)).pages) == 2
    wb = load_workbook(BytesIO(build_excel(bom, financial)))
    assert wb["Equipment BOM"]["A6"].data_type == "s"
    assert wb["Equipment BOM"]["A6"].value == "=2+0"
    assert wb["Inputs"]["C17"].value == "='Equipment BOM'!T156"


@pytest.mark.parametrize(
    "extension,signature,content_type",
    [("excel", b"PK", "spreadsheetml"), ("pdf", b"%PDF", "application/pdf")],
)
def test_export_endpoints(extension, signature, content_type):
    response = client.post(
        f"/api/export/{extension}", json=demo_assessment().model_dump()
    )
    assert response.status_code == 200
    assert response.content.startswith(signature)
    assert content_type in response.headers["content-type"]
    assert "attachment;" in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"
    demo = client.get(f"/api/export/{extension}")
    assert demo.status_code == 200
    assert "demo" in demo.headers["content-disposition"]
    if extension == "pdf":
        assert (
            "Synthetic demonstration"
            in PdfReader(BytesIO(demo.content)).pages[0].extract_text()
        )
    assert client.post(f"/api/export/{extension}", json={}).status_code == 422


def test_histogram_uses_all_scenarios(result):
    assert sum(bin_.count for bin_ in result.uncertainty.histogram) == 1000
    assert len(result.uncertainty.histogram) == 24


def test_large_constant_npv_histogram_is_finite():
    inputs = FinancialInputs(
        total_capex_usd=1000,
        plant_capacity_units_per_year=1e12,
        product_selling_price_usd_per_unit=1000,
    )
    risk = monte_carlo(inputs)
    assert sum(bin_.count for bin_ in risk.histogram) == 1000
    assert risk.p10_npv_usd == risk.p90_npv_usd


def test_scope_warnings_are_prominent_in_both_active_scenario_exports():
    payload = demo_assessment().model_dump(mode="json")
    payload["financial"]["discount_rate"] = 0.105
    payload["costing"]["equipment"] += [
        payload["costing"]["equipment"][1] | {"unit_id": "R1", "kind": "reactor"},
        {"unit_id": "Q1", "kind": "quoted", "purchased_cost_usd": 100, "installed_cost_usd": 200, "source": "Synthetic quote"},
    ]
    mixed = "MIXED EQUIPMENT SCOPE: retained equipment absent from native import."
    incomplete = "INCOMPLETE EQUIPMENT SCOPE: R-101 excluded from priced CAPEX."
    payload["engineering_notes"] = ["Operator note " + str(i) for i in range(8)] + [mixed + " Review retained equipment." * 20, incomplete + " Review excluded operations." * 20]
    calculated = client.post("/api/assessment", json=payload).json()
    excel = client.post("/api/export/excel", json=payload)
    pdf = client.post("/api/export/pdf", json=payload)
    assert excel.status_code == pdf.status_code == 200
    wb = load_workbook(BytesIO(excel.content))
    assert wb["Inputs"]["C7"].value == 0.105
    assert "NPV(" in wb["Summary"]["C6"].value
    notes = "\n".join(str(c.value) for row in wb["Notes"] for c in row)
    assert mixed in notes and incomplete in notes
    report = PdfReader(BytesIO(pdf.content))
    assert len(report.pages) == 2
    first = report.pages[0].extract_text()
    assert mixed in first and incomplete in first
    assert f"{calculated['financial']['npv_usd']:,.0f}" in first
