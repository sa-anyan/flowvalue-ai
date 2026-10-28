"""Persist synthetic fixtures and parsed results so failures can be independently reviewed."""

import asyncio
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
import pytest
import xlwt
from fastapi.testclient import TestClient

from app.agent.copilot_research import ResearchSettings
from app.agent.process_description import (
    MappingDraft,
    MentionedLink,
    compare_description,
    review_description,
)
from app.main import app, demo_assessment
from app.parsers.balance import parse_excel
from app.parsers.balance.validation import one

FIXTURES = Path(__file__).parent / "fixtures" / "balance"
OUTPUTS = Path(__file__).resolve().parents[2] / "outputs" / "balance-parser-tests"
OUTPUTS.mkdir(parents=True, exist_ok=True)


def workbook(name, rows, sheets=None, merged=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Streams"
    for row in rows:
        ws.append(row)
    for rng in merged:
        ws.merge_cells(rng)
    for title, data in (sheets or {}).items():
        sheet = wb.create_sheet(title)
        for row in data:
            sheet.append(row)
    data = BytesIO()
    wb.save(data)
    content = data.getvalue()
    (FIXTURES / name).write_bytes(content)
    return content


def parsed(name, content):
    result = parse_excel(content, name)
    (OUTPUTS / (name + ".json")).write_text(result.model_dump_json(indent=2))
    return result


def stream_rows():
    return [
        [
            "Stream ID",
            "Mass flow [kg/h]",
            "Temperature [C]",
            "Pressure [bara]",
            "Phase",
            "From",
            "To",
            "Boundary direction",
        ],
        ["F1", 100, 25, 1.2, "liquid", "SOURCE", "P-101", "inlet"],
        ["P1", 100, 30, 2, "liquid", "P-101", "SINK", "outlet"],
    ]


def test_simple_rows_source_lineage_and_zero():
    rows = stream_rows()
    rows[1][2] = 0
    r = parsed("simple.xlsx", workbook("simple.xlsx", rows))
    assert len(r.entities) == 2 and r.tables[0].orientation == "records_in_rows"
    v = one(r.entities[0], "mass_flow")
    assert v.normalized_value == 100 and v.cell == "B2" and v.header_cells == ["B1"]
    assert (
        v.workbook == "simple.xlsx"
        and v.sheet == "Streams"
        and v.original_unit == "kg/h"
    )
    assert one(r.entities[0], "temperature").normalized_value == 0
    assert all(
        c.status == "unresolved" for c in r.checks
    )  # no implicit complete/steady-state declaration


def test_transposed_units_column_and_merged_heading():
    rows = [
        ["Process stream conditions", None, None, None],
        ["Stream ID", "Units", "A1", "B1"],
        ["Mass flow", "kg/s", 1, 2],
        ["Temperature", "K", 300, 350],
        ["Pressure", "psia", 14.6959488, 29.3918976],
        ["Phase", None, "gas", "vapor"],
    ]
    r = parsed("columns.xlsx", workbook("columns.xlsx", rows, merged=["A1:D1"]))
    assert r.tables[0].orientation == "records_in_columns" and len(r.entities) == 2
    assert one(r.entities[0], "mass_flow").normalized_value == 3600
    assert one(r.entities[0], "temperature").normalized_value == pytest.approx(26.85)
    assert one(r.entities[0], "pressure").normalized_value == pytest.approx(
        1.01325, rel=1e-5
    )


def test_multisheet_multiple_tables_equipment_energy_and_connections():
    rows = stream_rows() + [[], ["Energy stream ID", "Heat added [kW]"], ["Q1", 5]]
    r = parsed(
        "mixed.xlsx",
        workbook(
            "mixed.xlsx",
            rows,
            {
                "Equipment": [
                    ["Tag", "Description", "Type", "Volume [m3]"],
                    ["P-101", "Feed pump", "pump", 2],
                ],
                "Extra": [["Balance basis", "steady state; complete boundary"]],
            },
        ),
    )
    assert {e.kind for e in r.entities} == {"material", "energy", "equipment"}
    assert len(r.tables) == 3
    assert (
        next(c for c in r.checks if c.name == "Overall mass balance").status
        == "verified"
    )
    assert (
        next(c for c in r.checks if c.name == "Energy balance").status == "unresolved"
    )
    assert not any(i.code == "unresolved_connection" for i in r.issues)


def test_component_balance_and_energy_requires_all_terms():
    rows = [
        ["Balance basis", "steady state; complete boundary"],
        ["Component balance basis", "nonreacting"],
        [
            "Energy balance basis",
            "steady state; complete boundary; kinetic and potential negligible; common enthalpy reference",
        ],
        ["Energy sign convention", "positive into boundary"],
        [],
        [
            "Stream ID",
            "Mass flow [kg/h]",
            "Boundary direction",
            "Water mass fraction [%]",
            "Enthalpy flow [kW]",
        ],
        ["F1", 100, "inlet", 100, 10],
        ["P1", 100, "outlet", 100, 15],
        [],
        [
            "Energy stream ID",
            "Heat added [kW]",
            "Heat removed [kW]",
            "Work in [kW]",
            "Work out [kW]",
        ],
        ["Q1", 5, 0, 0, 0],
    ]
    r = parsed("closure.xlsx", workbook("closure.xlsx", rows))
    for name in ("Overall mass balance", "Component mass: water", "Energy balance"):
        check = next(c for c in r.checks if c.name == name)
        assert check.status == "verified", check
        assert check.residual == pytest.approx(0)


def test_different_cases_not_summed_together_and_side_by_side():
    rows = [
        ["Case", "Stream ID", "Mass flow [kg/h]", None, "Stream ID", "Temperature [C]"],
        ["Full", "F1", 10, None, "X1", 25],
        ["Turn", "F1", 5, None, "X2", 30],
    ]
    r = parsed("cases.xlsx", workbook("cases.xlsx", rows))
    streams = [e for e in r.entities if e.identifier == "F1"]
    assert {e.case for e in streams} == {"Full", "Turn"}
    assert not any(i.code == "duplicate_identifier" for i in r.issues)
    assert len(r.tables) == 2


def test_second_unit_header_row_and_unit_normalization():
    rows = [
        [
            "Stream number",
            "Mass flow",
            "Temperature",
            "Density",
            "Molar flow",
            "Vapour fraction",
        ],
        [None, "lb/h", "F", "g/cm3", "mol/s", "%"],
        ["S1", 100, 32, 1, 2, 50],
    ]
    r = parsed("units.xlsx", workbook("units.xlsx", rows))
    e = r.entities[0]
    assert one(e, "mass_flow").normalized_value == pytest.approx(45.359237)
    assert one(e, "temperature").normalized_value == pytest.approx(0)
    assert one(e, "density").normalized_value == 1000
    assert one(e, "molar_flow").normalized_value == 7.2
    assert one(e, "vapor_fraction").normalized_value == 0.5


def test_invalid_missing_ambiguous_and_duplicate_data_are_retained():
    rows = [
        ["Stream ID", "Mass flow [kg/h]", "Temperature [K]", "Pressure [bar]"],
        ["S1", -1, 0, 2],
        ["S1", "1,234", None, 2],
        [None, 10, 300, 2],
    ]
    r = parsed("invalid-values.xlsx", workbook("invalid-values.xlsx", rows))
    assert len(r.entities) == 2
    assert {"duplicate_identifier", "missing_identifier", "unresolved_value"} <= {
        i.code for i in r.issues
    }
    assert one(r.entities[0], "mass_flow") is None
    assert one(r.entities[0], "pressure") is None  # absolute vs gauge unspecified


def test_conflicting_same_id_across_sheets_not_silently_chosen():
    r = parsed(
        "conflict.xlsx",
        workbook(
            "conflict.xlsx",
            [["Stream ID", "Mass flow [kg/h]"], ["S1", 10]],
            {"More": [["Stream ID", "Mass flow [kg/h]"], ["S1", 20]]},
        ),
    )
    assert any(i.code == "conflicting_values" for i in r.issues)
    assert all(c.status == "unresolved" for c in r.checks)


def with_cache(content, cell="B2", value="100"):
    out = BytesIO()
    with ZipFile(BytesIO(content)) as src, ZipFile(out, "w", ZIP_DEFLATED) as dest:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                from xml.etree import ElementTree as ET

                root = ET.fromstring(data)
                ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                node = root.find(f'.//x:c[@r="{cell}"]/x:v', ns)
                assert node is not None
                node.text = value
                data = ET.tostring(root)
            dest.writestr(item, data)
    return out.getvalue()


def test_formula_cache_missing_cache_error_and_external_formula_never_executed():
    content = workbook(
        "formula.xlsx",
        [
            ["Stream ID", "Mass flow [kg/h]", "Temperature [C]", "Pressure [bara]"],
            ["S1", "=50+50", "=1+2", 1],
            ["S2", 10, "#DIV/0!", 1],
        ],
    )
    content = with_cache(content)
    (FIXTURES / "formula.xlsx").write_bytes(content)
    r = parsed("formula.xlsx", content)
    v = one(r.entities[0], "mass_flow")
    assert (
        v.normalized_value == 100
        and v.formula == "=50+50"
        and v.cached
        and v.status == "inferred"
    )
    assert r.entities[0].values[1].normalized_value is None
    assert any(i.code == "unresolved_value" for i in r.issues)


def test_synthetic_provenance_never_becomes_verified():
    r = parsed(
        "synthetic.xlsx",
        workbook("synthetic.xlsx", [["SYNTHETIC STREAMS"], *stream_rows()]),
    )
    assert all(e.status == "inferred" for e in r.entities)
    assert all(v.status != "verified" for e in r.entities for v in e.values)
    assert all(c.status != "verified" for c in r.checks)


def test_xls_stored_results_are_read_and_marked_inferred():
    b = xlwt.Workbook()
    s = b.add_sheet("Streams")
    for i, row in enumerate(stream_rows()):
        for j, v in enumerate(row):
            s.write(i, j, v)
    f = BytesIO()
    b.save(f)
    content = f.getvalue()
    (FIXTURES / "legacy.xls").write_bytes(content)
    r = parsed("legacy.xls", content)
    assert len(r.entities) == 2 and one(r.entities[0], "mass_flow").status == "inferred"
    assert any(i.code == "workbook_warning" for i in r.issues)


@pytest.mark.parametrize(
    "name,content",
    [
        ("fake.xlsx", b"not excel"),
        ("fake.xls", b"<html/>"),
        ("unsupported.xlsb", b"PK"),
        ("empty.xlsx", b""),
    ],
)
def test_malformed_files_rejected_and_saved(name, content):
    (OUTPUTS / name).write_bytes(content)
    with pytest.raises(ValueError):
        parse_excel(content, name)


def test_zip_expansion_and_oversized_dimensions_rejected():
    b = BytesIO()
    with ZipFile(b, "w", ZIP_DEFLATED) as z:
        z.writestr("bomb.xml", b"0" * 3_000_000)
    (OUTPUTS / "expansion.xlsx").write_bytes(b.getvalue())
    with pytest.raises(ValueError, match="expansion"):
        parse_excel(b.getvalue(), "expansion.xlsx")
    w = openpyxl.Workbook()
    w.active["A10001"] = "x"
    out = BytesIO()
    w.save(out)
    (OUTPUTS / "dimensions.xlsx").write_bytes(out.getvalue())
    with pytest.raises(ValueError, match="limit"):
        parse_excel(out.getvalue(), "dimensions.xlsx")


def test_description_grounding_contradictions_and_no_mutation():
    content = workbook(
        "description.xlsx",
        [["Stream ID", "From", "To", "Mass flow [kg/h]"], ["S1", "R-101", "R-102", 10]],
        {
            "Equipment": [
                ["Tag", "Description"],
                ["R-101", "Mixer"],
                ["R-102", "Reactor"],
            ]
        },
    )
    r = parsed("description.xlsx", content)
    before = r.model_dump_json()
    text = "R-102 feeds R-101. Missing P-999."
    result = compare_description(
        r,
        text,
        MappingDraft(
            links=[
                MentionedLink(
                    source_tag="R-102",
                    destination_tag="R-101",
                    quote="R-102 feeds R-101",
                )
            ],
            unresolved_excerpts=[],
        ),
        "mock",
    )
    assert {"description_contradiction", "description_unknown_tag"} <= {
        i.code for i in result.findings
    }
    assert r.model_dump_json() == before
    with pytest.raises(ValueError):
        compare_description(
            r,
            text,
            MappingDraft(
                links=[
                    MentionedLink(
                        source_tag="EVIL", destination_tag="R-101", quote="made up"
                    )
                ],
                unresolved_excerpts=[],
            ),
        )
    (OUTPUTS / "description-review.json").write_text(result.model_dump_json(indent=2))


def test_ai_uses_bounded_structured_mapping_no_tools_and_rejects_incomplete():
    r = parsed(
        "ai.xlsx",
        workbook(
            "ai.xlsx",
            stream_rows(),
            {"Equipment": [["Tag", "Description"], ["P-101", "Feed pump"]]},
        ),
    )
    calls = []

    class Client:
        responses = None

        def __init__(self):
            self.responses = self

        async def parse(self, **kw):
            calls.append(kw)
            return SimpleNamespace(
                status="completed",
                output_parsed=MappingDraft(links=[], unresolved_excerpts=["raw water"]),
            )

    result = asyncio.run(
        review_description(
            r,
            "raw water enters a filter",
            ResearchSettings(openai_api_key="mock"),
            Client(),
        )
    )
    assert result.questions and calls[0]["store"] is False and "tools" not in calls[0]
    assert "100" not in calls[0]["input"]  # numerical balance values never sent
    assert calls[0]["text_format"] == MappingDraft


def test_upload_api_and_description_basic_check(monkeypatch):
    monkeypatch.setenv("REQUIRE_AUTH", "false")
    c = TestClient(app)
    content = workbook("api.xlsx", stream_rows())
    (OUTPUTS / "api-bad.xlsx").write_bytes(b"no")
    (OUTPUTS / "api-huge.xlsx").write_bytes(b"0" * (4 * 1024 * 1024 + 1))
    response = c.post(
        "/api/streams/parse/excel",
        files={
            "file": (
                "api.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200, response.text
    result = c.post(
        "/api/process-description/check",
        json={"report": response.json(), "description": "P-999 receives water"},
    )
    assert (
        result.status_code == 200
        and result.json()["findings"][0]["code"] == "description_unknown_tag"
    )
    assert (
        c.post(
            "/api/streams/parse/excel", files={"file": ("bad.xlsx", b"no")}
        ).status_code
        == 422
    )
    assert (
        c.post(
            "/api/streams/parse/excel",
            files={"file": ("huge.xlsx", b"0" * (4 * 1024 * 1024 + 1))},
        ).status_code
        == 413
    )


def test_saved_attachment_round_trips_through_existing_project_endpoint(monkeypatch):
    from app.auth import authenticated_user

    monkeypatch.setenv("REQUIRE_AUTH", "false")

    async def user():
        return {"id": "owner"}

    writes = []

    async def rest(request, method, path, *, body=None):
        writes.append(body)
        return {"revision": 1}

    app.dependency_overrides[authenticated_user] = user
    monkeypatch.setattr("app.main.project_rest", rest)
    r = parsed("saved.xlsx", workbook("saved.xlsx", stream_rows()))
    try:
        response = TestClient(app).post(
            "/api/projects",
            json={
                "name": "Test",
                "project": demo_assessment().model_dump(mode="json"),
                "balance_import": {"report": r.model_dump(mode="json")},
            },
        )
        assert response.status_code == 200, response.text
        assert writes[0]["p_payload"]["balance_import"]["report"]["sha256"] == r.sha256
    finally:
        app.dependency_overrides.clear()


def test_long_property_table_and_unmapped_property_retained():
    rows = [
        ["Stream ID", "Property", "Units", "Value"],
        ["S1", "Mass flow", "kg/s", 2],
        ["S1", "Temperature", "K", 300],
        ["S1", "Custom index", "index", 7],
        ["S2", "Mass flow", "kg/h", 5],
    ]
    r = parsed("long-format.xlsx", workbook("long-format.xlsx", rows))
    assert len(r.entities) == 2
    assert one(r.entities[0], "mass_flow").normalized_value == 7200
    assert one(r.entities[0], "temperature").cell == "D3"
    assert (
        r.entities[0].values[2].original_value == 7
        and r.entities[0].values[2].status == "unresolved"
    )


def test_unsupported_units_invalid_phase_and_pressure_basis_are_explicit():
    r = parsed(
        "unknown-units.xlsx",
        workbook(
            "unknown-units.xlsx",
            [
                [
                    "Stream ID",
                    "Mass flow [widgets/day]",
                    "Absolute pressure [kPa]",
                    "Phase",
                ],
                ["S1", 10, 101.325, "plasma"],
            ],
        ),
    )
    v = r.entities[0].values[0]
    assert v.original_unit == "widgets/day" and v.normalized_value is None
    assert one(r.entities[0], "pressure").normalized_value == pytest.approx(1.01325)
    assert one(r.entities[0], "phase") is None and r.entities[0].status == "unresolved"


def test_conflicting_declarations_block_validation():
    rows = [["Balance basis", "steady state; complete boundary"], [], *stream_rows()]
    r = parsed(
        "declaration-conflict.xlsx",
        workbook(
            "declaration-conflict.xlsx",
            rows,
            {"Other": [["Balance basis", "transient; incomplete boundary"]]},
        ),
    )
    assert any(i.code == "conflicting_declarations" for i in r.issues)
    assert all(c.status != "verified" for c in r.checks)


def test_excel_percent_storage_is_not_divided_twice():
    w = openpyxl.Workbook()
    s = w.active
    s.append(["Stream ID", "Water mass fraction [%]"])
    s.append(["S1", 0.8])
    s["B2"].number_format = "0.0%"
    buf = BytesIO()
    w.save(buf)
    (FIXTURES / "percent.xlsx").write_bytes(buf.getvalue())
    r = parsed("percent.xlsx", buf.getvalue())
    assert r.entities[0].values[0].normalized_value == 0.8


def test_missing_zero_work_terms_prevents_energy_closure():
    r = parse_excel((FIXTURES / "closure.xlsx").read_bytes(), "closure.xlsx")
    energy = next(e for e in r.entities if e.kind == "energy")
    energy.values = [v for v in energy.values if v.property == "heat_added"]
    r.checks = []
    from app.parsers.balance.validation import validate

    validate(r)
    assert (
        next(c for c in r.checks if c.name == "Energy balance").status == "unresolved"
    )
    (OUTPUTS / "missing-zero-work-terms.json").write_text(r.model_dump_json(indent=2))


def test_api_requires_auth_and_ai_disabled_is_explicit(monkeypatch):
    monkeypatch.setenv("REQUIRE_AUTH", "true")
    c = TestClient(app)
    assert (
        c.post(
            "/api/streams/parse/excel",
            files={"file": ("simple.xlsx", (FIXTURES / "simple.xlsx").read_bytes())},
        ).status_code
        == 401
    )
    monkeypatch.setenv("REQUIRE_AUTH", "false")
    monkeypatch.setenv("AI_ENABLED", "false")
    report = parse_excel(
        (FIXTURES / "simple.xlsx").read_bytes(), "simple.xlsx"
    ).model_dump(mode="json")
    r = c.post(
        "/api/process-description/review",
        json={"report": report, "description": "Raw water enters a pump."},
    )
    assert r.status_code == 503 and "disabled" in r.text


def test_equipment_inlet_outlet_conflicts_and_unknowns_are_visible():
    r = parsed(
        "equipment-links.xlsx",
        workbook(
            "equipment-links.xlsx",
            [["Stream ID", "From", "To"], ["S1", "R-101", "R-102"]],
            {
                "Equipment": [
                    ["Tag", "Inlet streams", "Outlet streams"],
                    ["R-101", "S1", "MISSING"],
                    ["R-102", "S1", None],
                ]
            },
        ),
    )
    assert {"connection_conflict", "unresolved_equipment_stream"} <= {
        i.code for i in r.issues
    }


def test_xml_entities_and_macro_payloads_rejected_without_execution():
    base = (FIXTURES / "simple.xlsx").read_bytes()
    for name, payload, expected in [
        (
            "unsafe-xml.xlsx",
            b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x>&e;</x>',
            "DTD",
        ),
        ("macro.xlsx", b"not executable", "Macro"),
    ]:
        out = BytesIO()
        with ZipFile(BytesIO(base)) as src, ZipFile(out, "w", ZIP_DEFLATED) as z:
            for item in src.infolist():
                z.writestr(item, src.read(item.filename))
            z.writestr(
                "extra.xml" if name == "unsafe-xml.xlsx" else "xl/vbaProject.bin",
                payload,
            )
        (OUTPUTS / name).write_bytes(out.getvalue())
        with pytest.raises(ValueError, match=expected):
            parse_excel(out.getvalue(), name)


def test_raw_request_limits_apply_before_multipart_and_json_parsing(monkeypatch):
    monkeypatch.setenv("REQUIRE_AUTH", "false")
    c = TestClient(app)
    for route in (
        "/api/streams/parse/excel",
        "/api/process-description/check",
        "/api/process-description/review",
    ):
        response = c.post(
            route,
            content=b"{}",
            headers={"content-length": "99999999", "content-type": "application/json"},
        )
        assert response.status_code == 413
