from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError
from pypdf import PdfReader

from app.engine.costing import CostingRequest, build_bom
from app.engine.indices import IndexRecord, index_catalog
from app.financial.dcf_model import (
    FinancialInputs,
    MonteCarloSettings,
    ProductionIncentive,
    calculate_dcf,
    monte_carlo,
)
from app.financial.sensitivity import tornado
from app.integrations.libraries import OpenPyTEARequest, openpytea_equipment
from app.main import app, assessment, demo_assessment
from app.parsers.dwsim_parser import parse_dwxml
from app.reports.excel_builder import build_excel
from app.reports.pdf_builder import build_pdf

client = TestClient(app)
FIXTURE = Path(__file__).parent / "fixtures/synthetic.dwxml"


def financial(**changes):
    return FinancialInputs(
        total_capex_usd=1000,
        plant_capacity_units_per_year=100,
        product_selling_price_usd_per_unit=2,
        operating_years=3,
        discount_rate=0,
        tax_rate=0.2,
        **changes,
    )


def policy(**changes):
    return ProductionIncentive(
        enabled=True,
        eligibility_confirmed=True,
        basis="Synthetic scenario, not a claim of legal eligibility",
        **changes,
    )


def test_dwxml_native_duplicate_properties_si_units_and_mapping():
    result = parse_dwxml(FIXTURE.read_bytes())
    assert result.streams[0].mass_flow_kg_h == 36000
    assert result.streams[0].temperature_c == pytest.approx(76.85)
    assert result.streams[0].pressure_bar_abs == 3
    assert result.streams[0].phase == "liquid"
    assert result.streams[0].destination_unit == "E-101"
    assert result.streams[1].source_unit == "E-101"
    assert result.units[0].x == 220
    assert result.equipment[0]["duty_kw"] == 1000
    assert result.equipment[0]["design_pressure_barg"] == pytest.approx(1.98675)
    assert result.units[1].equipment is None
    assert "volume_m3" in result.units[1].missing_inputs
    assert (
        build_bom(CostingRequest(equipment=result.equipment)).items[0].sizing["area_m2"]
        == 100
    )


@pytest.mark.parametrize(
    "content",
    [
        b'<!DOCTYPE root [<!ENTITY xxe SYSTEM "file:///etc/passwd">]><DWSIM_Simulation_Data>&xxe;</DWSIM_Simulation_Data>',
        b"<broken>",
        b"<anything/>",
        b"<DWSIM_Simulation_Data/>",
        b"x" * (5 * 1024 * 1024 + 1),
    ],
)
def test_unsafe_malformed_or_empty_xml_rejected(content):
    with pytest.raises(ValueError):
        parse_dwxml(content)


def test_dwxml_namespaces_and_uncalculated_unit_review():
    data = FIXTURE.read_bytes().replace(
        b"<DWSIM_Simulation_Data>", b'<DWSIM_Simulation_Data xmlns="urn:dwsim">'
    )
    data = data.replace(
        b"<Calculated>true</Calculated><Q>", b"<Calculated>false</Calculated><Q>"
    )
    result = parse_dwxml(data)
    assert len(result.streams) == 2
    assert not result.equipment
    assert "Re-converge" in result.units[0].missing_inputs[0]


def test_xml_upload_and_size_limit():
    r = client.post(
        "/api/streams/parse/dwxml",
        files={"file": ("synthetic.dwxml", FIXTURE.read_bytes(), "application/xml")},
    )
    assert r.status_code == 200
    assert len(r.json()["equipment"]) == 1
    assert (
        client.post(
            "/api/streams/parse/dwxml",
            files={"file": ("large.dwxml", b"x" * (5 * 1024 * 1024 + 1))},
        ).status_code
        == 413
    )


def test_incentives_require_basis_and_eligibility():
    with pytest.raises(ValidationError):
        ProductionIncentive(enabled=True)
    with pytest.raises(ValidationError):
        ProductionIncentive(enabled=True, eligibility_confirmed=True)


def test_cash_credit_start_duration_dcf_and_uncertainty():
    base = calculate_dcf(financial())
    credited = financial(
        incentive=policy(
            rate_usd_per_unit=1, eligible_fraction=0.5, start_year=2, duration_years=1
        )
    )
    result = calculate_dcf(credited)
    assert [r.credit_cash_usd for r in result.cash_flows] == [0, 0, 50, 0]
    assert result.npv_usd - base.npv_usd == 50
    assert result.lcop_usd_per_unit == base.lcop_usd_per_unit
    assert (
        monte_carlo(credited, MonteCarloSettings(variation_fraction=0)).p50_npv_usd
        == result.npv_usd
    )


def test_nonrefundable_tax_credit_does_not_create_refund():
    inputs = financial(incentive=policy(rate_usd_per_unit=100, treatment="tax_offset"))
    result = calculate_dcf(inputs)
    assert all(r.taxes_usd == 0 for r in result.cash_flows)
    assert all(r.credit_cash_usd == 0 for r in result.cash_flows)
    assert result.cash_flows[1].credit_generated_usd == 10000
    assert result.cash_flows[1].credit_tax_offset_usd == 0  # loss after depreciation
    assert (
        monte_carlo(inputs, MonteCarloSettings(variation_fraction=0)).p50_npv_usd
        == result.npv_usd
    )


def test_tax_offset_capped_by_tax_liability():
    inputs = financial(
        incentive=policy(rate_usd_per_unit=100, treatment="tax_offset"),
        depreciation_years=1,
    )
    result = calculate_dcf(inputs)
    assert result.cash_flows[2].credit_tax_offset_usd == 40
    assert result.cash_flows[2].free_cash_flow_usd == 200


def test_expired_credit_with_negative_cashflows_does_not_claim_unique_irr():
    inputs = financial(
        fixed_opex_usd_per_year=300,
        incentive=policy(rate_usd_per_unit=10, duration_years=1),
    )
    result = calculate_dcf(inputs)
    assert result.cash_flows[1].free_cash_flow_usd > 0
    assert result.cash_flows[2].free_cash_flow_usd < 0
    assert result.irr is None


def test_tornado_sorted_and_fixed_variable_subset_consistent():
    inputs = financial(
        variable_opex_usd_per_year=100,
        feedstock_opex_usd_per_year=60,
        utility_opex_usd_per_year=40,
    )
    result = tornado(inputs)
    assert len(result["rows"]) == 6
    assert result["rows"] == sorted(
        result["rows"], key=lambda r: r["span_usd"], reverse=True
    )
    utility = next(r for r in result["rows"] if r["label"] == "Utilities")
    changed = inputs.model_dump()
    changed.update(utility_opex_usd_per_year=48, variable_opex_usd_per_year=108)
    assert (
        utility["high_input_npv_usd"]
        == calculate_dcf(FinancialInputs(**changed)).npv_usd
    )


def test_real_openpytea_adapter_and_quoted_equipment_export():
    result = openpytea_equipment(
        OpenPyTEARequest(
            name="C1",
            parameter=5000,
            category="Compressors, fans, & Blowers",
            equipment_type="Compressor, centrifugal",
        )
    )
    assert result["version"] == "3.2.1"
    assert result["equipment"]["purchased_cost_usd"] > 0
    bom = build_bom(CostingRequest(equipment=[result["equipment"]]))
    fin = calculate_dcf(
        financial().model_copy(
            update={"total_capex_usd": bom.total_bare_module_cost_usd}
        )
    )
    wb = load_workbook(BytesIO(build_excel(bom, fin)))
    assert (
        wb["Equipment BOM"]["N6"].data_type == "n"
    )  # external quote is an editable input
    assert wb["Equipment BOM"]["O6"].data_type == "f"
    assert wb["Equipment BOM"]["R6"].data_type == "f"
    assert (
        len(PdfReader(BytesIO(build_pdf(bom, fin, monte_carlo(fin.inputs)))).pages) == 2
    )


def test_biosteam_snapshot_becomes_valid_assessment():
    payload = {
        "schema_version": "flowvalue-biosteam-v1",
        "library_version": "2.52.0",
        "reference_cepci": 800,
        "provenance": "Synthetic exported snapshot, not an actual simulation",
        "equipment": [
            {"unit_id": "BIO1", "purchased_cost_usd": 1000, "installed_cost_usd": 3000}
        ],
        "fixed_opex_usd_per_year": 10,
        "variable_opex_usd_per_year": 100,
        "feedstock_opex_usd_per_year": 60,
        "utility_opex_usd_per_year": 40,
        "production_units_per_year": 100,
        "selling_price_usd_per_unit": 2,
    }
    r = client.post("/api/integrations/biosteam/import", json=payload)
    assert r.status_code == 200
    evaluated = client.post("/api/assessment", json=r.json())
    assert evaluated.status_code == 200
    assert evaluated.json()["bom"]["total_bare_module_cost_usd"] == 3000


def test_policy_excel_formulas_and_two_page_pdf():
    request = demo_assessment()
    request.financial.incentive = policy(
        name="Synthetic 45V scenario", rate_usd_per_unit=0.5
    )
    result = assessment(request)
    wb = load_workbook(BytesIO(build_excel(result.bom, result.financial)))
    assert wb["Inputs"]["C27"].value == 1
    assert wb["Inputs"]["C28"].value == 0.5
    for row in wb["Cash Flow"].iter_rows(min_row=6, max_row=106, max_col=22):
        assert all(c.data_type == "f" for c in row)
    assert "U7" in wb["Cash Flow"]["L7"].value
    assert "V7" in wb["Cash Flow"]["I7"].value
    reader = PdfReader(
        BytesIO(build_pdf(result.bom, result.financial, result.uncertainty))
    )
    assert len(reader.pages) == 2
    assert "Synthetic 45V scenario" in reader.pages[1].extract_text()


def test_catalog_does_not_claim_live_verified_index():
    catalog = index_catalog()
    assert catalog["live_feed"] is False
    record = catalog["records"][0] | {"retrieved_on": date.today() + timedelta(days=1)}
    with pytest.raises(ValidationError):
        IndexRecord.model_validate(record)


def test_projects_require_authenticated_user():
    assert client.get("/api/projects").status_code == 401
    assert client.post("/api/projects", json={}).status_code == 401
    assert (
        client.get("/api/projects/10000000-0000-4000-8000-000000000001").status_code
        == 401
    )


def test_required_auth_fail_closed(monkeypatch):
    monkeypatch.setenv("REQUIRE_AUTH", "true")
    assert client.get("/api/demo").status_code == 401
    assert client.get("/api/config").status_code == 200
    assert client.get("/health").status_code == 200
    assert client.get(
        "/api/demo", headers={"Authorization": "Bearer invalid"}
    ).status_code in (401, 503)


def test_incomplete_scope_survives_assessment_and_exports():
    request = demo_assessment()
    note = "INCOMPLETE EQUIPMENT SCOPE: R-101 excluded from priced CAPEX."
    request.engineering_notes = [note]
    result = assessment(request)
    assert note in result.bom.assumptions
    assert note in result.financial.warnings
    wb = load_workbook(BytesIO(build_excel(result.bom, result.financial)))
    assert any(note in str(c.value) for row in wb["Notes"] for c in row)
    reader = PdfReader(
        BytesIO(build_pdf(result.bom, result.financial, result.uncertainty))
    )
    assert len(reader.pages) == 2
    assert note in reader.pages[0].extract_text()
