"""Independent expectations for source lineage, shared registers and supported layouts."""

from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
import pytest
from defusedxml.ElementTree import fromstring

from app.parsers.balance.models import BalanceReport
from app.parsers.balance.workbook import read_book
from app.parsers.workbook_pipeline import parse_workbook

FIXTURES = Path(__file__).parent / "fixtures/parser_v3"


def workbook(name, sheets, merges=()):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    for title, ref in merges:
        book[title].merge_cells(ref)
    output = BytesIO()
    book.save(output)
    raw = output.getvalue()
    # Retain exactly the workbook parsed by each regression for manual inspection.
    (FIXTURES / name).write_bytes(raw)
    return raw


def materials(report, sheet=None):
    return [
        e
        for e in report.entities
        if e.kind == "material" and (sheet is None or e.id_evidence.sheet == sheet)
    ]


def value(entity, prop, component=None):
    return next(
        v for v in entity.values if v.property == prop and v.component == component
    )


def assert_graph(report):
    entities = {e.key: e for e in report.entities}
    assert len(entities) == len(report.entities)
    assert all(e.table in {t.id for t in report.tables} for e in report.entities)
    for e in report.entities:
        for key in e.stream_keys:
            assert entities[key].kind == "material"
        if e.kind == "equipment":
            assert not any(v.property.startswith("component_") for v in e.values)
    for edge in report.connections:
        assert edge.stream_key in entities
        assert (
            edge.equipment_key is None
            or entities[edge.equipment_key].kind == "equipment"
        )
    BalanceReport.model_validate_json(report.model_dump_json())


def test_team8_independent_source_expectations():
    report = parse_workbook(
        (FIXTURES / "team8-reference.xlsm").read_bytes(), "team8-reference.xlsm"
    )
    streams = materials(report, "Mass and Energy Balance")
    assert [s.identifier for s in streams] == [str(i) for i in range(1, 15)]
    assert len(materials(report, "Energy calculation")) == 18
    assert len(materials(report)) == 32  # template calculation headings are not streams
    assert {e.identifier for e in report.entities if e.kind == "equipment"} == {
        "B-001",
        "P-001",
        "R-001",
        "S-001",
        "T-001",
        "T-002",
        "T-003",
    }
    expected = [
        (6.25, 0.00367, 1, 25, 34.5, {"Water": 0.125, "Sulphuric Acid": 6.125}),
        (6.25, 0.00367, 3, 25, 34.5, {"Water": 0.125, "Sulphuric Acid": 6.125}),
        (
            125127.5,
            125.1275,
            1,
            30,
            125.4,
            {"Water": 125000, "Lignin": 125, "Ammonia": 2.5},
        ),
        (
            125127.5,
            125.1275,
            2,
            30,
            125.4,
            {"Water": 125000, "Lignin": 125, "Ammonia": 2.5},
        ),
        (
            125133.75,
            125.13117,
            1,
            30,
            125.4,
            {
                "Water": 125000.125,
                "Lignin": 125,
                "Sulphuric Acid": 6.125,
                "Ammonia": 2.5,
            },
        ),
        (66.66, 0.0542, 1, 25, 87.5, {"Water": 33.33, "Sucrose": 33.33}),
        (66.66, 0.0542, 2, 25, 87.5, {"Water": 33.33, "Sucrose": 33.33}),
        (1.2, 0.00299, 1, 25, 83.5, {"Water": 0.94, "NaOH": 0.26}),
        (1.2, 0.00299, 2, 25, 83.5, {"Water": 0.94, "NaOH": 0.26}),
        (74, 62.5, 1.5, 25, 25.25, {"Air": 74}),
        (95.5, 80.6, 1, 30, 30.3, {"CO2": 32.7, "Air": 62.8}),
        (
            125927.625,
            125.88489,
            1.1,
            30,
            125.4,
            {"Water": 125900.125, "Lignin": 25, "Ammonia": 2.5, "Biomass": 100},
        ),
        (
            125027.625,
            125.02789,
            1,
            30,
            125.4,
            {"Water": 125000.125, "Lignin": 25, "Ammonia": 2.5},
        ),
        (900, 0.857, 1, 30, 125.4, {"Water": 900, "Biomass": 100}),
    ]
    expected_volumes = [
        {"Water": 0.00039, "Sulphuric Acid": 0.00328},
        {"Water": 0.00039, "Sulphuric Acid": 0.00328},
        {"Water": 125, "Lignin": 0.125, "Ammonia": 0.0025},
        {"Water": 125, "Lignin": 0.125, "Ammonia": 0.0025},
        {
            "Water": 125.00039,
            "Lignin": 0.125,
            "Sulphuric Acid": 0.00328,
            "Ammonia": 0.0025,
        },
        {"Water": 0.0271, "Sucrose": 0.0271},
        {"Water": 0.0271, "Sucrose": 0.0271},
        {"Water": 0.00234, "NaOH": 0.00065},
        {"Water": 0.00234, "NaOH": 0.00065},
        {"Air": 62.5},
        {"CO2": 27.6, "Air": 53},
        {"Water": 125.85739, "Lignin": 0.025, "Ammonia": 0.0025, "Biomass": "-"},
        {"Water": 125.00039, "Lignin": 0.025, "Ammonia": 0.0025},
        {"Water": 0.857, "Biomass": "-"},
    ]
    for stream, volumes in zip(streams, expected_volumes, strict=True):
        assert {
            v.component: v.original_value
            for v in stream.values
            if v.property == "component_volume_flow"
        } == volumes
    cols = [
        "O",
        "Y",
        "AI",
        "AS",
        "BC",
        "BM",
        "BW",
        "CG",
        "CQ",
        "DA",
        "O",
        "Y",
        "AI",
        "AS",
    ]
    for index, (stream, spec) in enumerate(zip(streams, expected, strict=True)):
        mass, volume, pressure, temperature, enthalpy, components = spec
        for prop, number in [
            ("mass_flow", mass),
            ("volume_flow", volume),
            ("pressure", pressure),
            ("temperature", temperature),
            ("specific_enthalpy", enthalpy),
        ]:
            assert value(stream, prop).original_value == pytest.approx(number)
        assert {
            v.component: v.original_value
            for v in stream.values
            if v.property == "component_mass_flow"
        } == components
        assert stream.id_evidence.cell == f"{cols[index]}{8 if index < 10 else 25}"
        assert (
            value(stream, "temperature").cell
            == f"{cols[index]}{22 if index < 10 else 39}"
        )
        assert value(stream, "pressure").normalized_unit == "bar_abs"
    assert value(streams[11], "component_mass_flow", "Biomass").cell == "AD36"
    assert value(streams[13], "component_mass_flow", "Biomass").cell == "AX36"
    mismatches = [i for i in report.issues if i.code == "component_total_mismatch"]
    assert len(mismatches) == 2
    assert all(
        any(f"stream {sid} " in i.message for i in mismatches) for sid in ["12", "14"]
    )
    assert not any(
        edge.equipment_key
        for edge in report.connections
        if edge.stream_key in {streams[12].key, streams[13].key}
    )
    with ZipFile(FIXTURES / "team8-reference.xlsm") as archive:
        cells = {
            n.attrib.get("r"): n
            for n in fromstring(archive.read("xl/worksheets/sheet1.xml")).iter()
            if n.tag.endswith("}c")
        }
        formulas = {
            coord: next(c.text for c in cells[coord] if c.tag.endswith("}f"))
            for coord in ["AD37", "AX37"]
        }
        assert formulas == {"AD37": "SUM(AD28:AH35)", "AX37": "SUM(AX28:BB35)"}
    assert_graph(report)


def test_mixed_layout_blank_description_merged_ids_and_exact_component_names():
    raw = workbook(
        "mixed-merged.xlsm",
        {
            "Mixed": [
                ["Stream Number", "S-01", None, "S-02", None],
                ["Description", None, None, "Water to tank T-001"],
                ["Component", "m3/h", "kg/h", "m3/h", "kg/h"],
                ["Na+", 0.1, 100, 0.2, 200],
                ["α-amylase", 0.1, 2, 0.2, 4],
                ["Total", 0.2, 102, 0.4, 204],
                [],
                ["Stream ID", "Mass flow (kg/s)", "Source", "Destination"],
                ["OUT-A", 2, "T-001", "outside"],
            ]
        },
        [("Mixed", "B1:C1"), ("Mixed", "D1:E1")],
    )
    report = parse_workbook(raw, "mixed-merged.xlsm")
    streams = materials(report)
    assert {e.identifier for e in streams} == {"S-01", "S-02", "OUT-A"}
    first = next(e for e in streams if e.identifier == "S-01")
    assert value(first, "component_mass_flow", "Na+").original_value == 100
    assert value(first, "component_mass_flow", "α-amylase").original_value == 2
    assert first.id_evidence.cell == "B1"
    assert_graph(report)


def test_scopes_repeated_ids_same_tag_and_conflicting_endpoint_evidence():
    rows = [
        [
            "Stream ID",
            "Source",
            "Destination",
            "Description",
            "Mass flow (kg/h)",
            "Case",
        ],
        ["A", "T-001", "P-001", "Water from tank T-002 to pump P-001", 100, "base"],
        ["A", "T-001", "P-001", "Water feed", 110, "upset"],
    ]
    report = parse_workbook(
        workbook("scopes.xlsx", {"Plant A": rows, "Plant B": rows}), "scopes.xlsx"
    )
    assert len(materials(report)) == 4
    assert len({s.key for s in materials(report)}) == 4
    assert any(i.code == "conflicting_connection" for i in report.issues)
    assert (
        len(
            [
                e
                for e in report.entities
                if e.kind == "equipment" and e.identifier == "T-001"
            ]
        )
        == 4
    )
    assert_graph(report)


def test_missing_component_unknown_unit_and_invalid_number_are_not_filled():
    report = parse_workbook(
        workbook(
            "unknown-units.xlsx",
            {
                "Data": [
                    ["Stream Number", "X1"],
                    ["Description", None],
                    ["Component", "furlong/h"],
                    [None, 4],
                    ["Water", "bad"],
                    ["Total", 5],
                ]
            },
        ),
        "unknown-units.xlsx",
    )
    stream = materials(report)[0]
    assert any(i.code == "missing_component" for i in report.issues)
    assert any(
        v.original_value == 4 and v.normalized_value is None for v in stream.values
    )
    assert any(
        v.original_value == "bad" and v.status == "unresolved" for v in stream.values
    )
    assert_graph(report)


def test_one_open_with_cached_formula_error_and_provenance(monkeypatch):
    raw = workbook(
        "formula-base.xlsm",
        {"Data": [["Stream ID", "Mass flow (kg/h)"], ["F-1", "=1+2"], ["F-2", "=1/0"]]},
    )
    buffer = BytesIO()
    with ZipFile(BytesIO(raw)) as source, ZipFile(buffer, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b"<v />", b"<v></v>").replace(
                    b"<v/>", b"<v></v>"
                )
                content = content.replace(b"<f>1+2</f><v></v>", b"<f>1+2</f><v>3</v>")
                content = content.replace(b'<c r="B3"', b'<c t="e" r="B3"').replace(
                    b"<f>1/0</f><v></v>", b"<f>1/0</f><v>#DIV/0!</v>"
                )
            target.writestr(item, content)
    raw = buffer.getvalue()
    (FIXTURES / "formula-cache.xlsm").write_bytes(raw)
    original = openpyxl.load_workbook
    calls = []

    def load(*args, **kwargs):
        calls.append(kwargs)
        return original(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", load)
    report = parse_workbook(raw, "formula-cache.xlsm")
    assert len(calls) == 1 and calls[0]["read_only"] and not calls[0]["keep_links"]
    streams = materials(report)
    assert value(streams[0], "mass_flow").original_value == 3
    assert value(streams[0], "mass_flow").formula == "=1+2"
    assert value(streams[1], "mass_flow").original_value == "#DIV/0!"
    assert value(streams[1], "mass_flow").status == "unresolved"


def test_hidden_rows_and_columns_are_not_discarded():
    raw = workbook(
        "hidden.xlsx", {"Data": [["Stream ID", "Mass flow (kg/h)"], ["S1", 10]]}
    )
    book = openpyxl.load_workbook(BytesIO(raw))
    book.active.row_dimensions[2].hidden = True
    book.active.column_dimensions["B"].hidden = True
    output = BytesIO()
    book.save(output)
    raw = output.getvalue()
    (FIXTURES / "hidden.xlsx").write_bytes(raw)
    assert (
        value(
            materials(parse_workbook(raw, "hidden.xlsx"))[0], "mass_flow"
        ).original_value
        == 10
    )


def test_format_and_resource_safeguards():
    for filename, content in [
        ("bad.xlsm", b"bad"),
        ("unsupported.xlsb", b"PKbad"),
        ("large.xlsm", b"x" * (4 * 1024 * 1024 + 1)),
    ]:
        with pytest.raises(ValueError):
            read_book(content, filename)


def test_numeric_description_block_does_not_hide_vertical_block():
    report = parse_workbook(
        workbook(
            "mixed-numeric.xlsx",
            {
                "Data": [
                    ["Stream Number", 1],
                    ["Description", "Water to tank T-101"],
                    ["Component", "kg/h"],
                    ["Water", 10],
                    [],
                    ["Stream ID", "Mass flow (kg/h)", "Source", "Destination"],
                    ["OUT-2", 10, "T-101", "outside"],
                ]
            },
        ),
        "mixed-numeric.xlsx",
    )
    assert {s.identifier for s in materials(report)} == {"1", "OUT-2"}
    assert_graph(report)


def test_duplicate_identifier_records_are_retained_and_flagged():
    report = parse_workbook(
        workbook(
            "duplicate.xlsx",
            {
                "Data": [
                    ["Stream ID", "Mass flow (kg/h)"],
                    ["S1", 10],
                    ["S1", 20],
                ]
            },
        ),
        "duplicate.xlsx",
    )
    assert len(materials(report)) == 2
    assert {value(s, "mass_flow").original_value for s in materials(report)} == {10, 20}
    assert any(i.code == "duplicate_identifier" for i in report.issues)
    assert_graph(report)


def test_missing_units_preserve_components_and_streams_without_description():
    report = parse_workbook(
        workbook(
            "missing-units.xlsx",
            {
                "Data": [
                    ["Stream Number", "FEED-A"],
                    ["Component", None],
                    ["Water", 10],
                    ["Lignin", 5],
                    ["Total", 15],
                ]
            },
        ),
        "missing-units.xlsx",
    )
    stream = materials(report)[0]
    assert stream.identifier == "FEED-A"
    assert {(v.component, v.original_value) for v in stream.values if v.component} == {
        ("Water", 10),
        ("Lignin", 5),
    }
    assert all(v.normalized_value is None for v in stream.values)
    assert any(i.code == "unresolved_value" for i in report.issues)


def test_long_form_case_column_does_not_merge_cases():
    report = parse_workbook(
        workbook(
            "long-cases.xlsx",
            {
                "Data": [
                    ["Stream ID", "Property", "Value", "Unit", "Case"],
                    ["FEED", "Mass flow", 10, "kg/h", "base"],
                    ["FEED", "Temperature", 25, "C", "base"],
                    ["FEED", "Mass flow", 20, "kg/h", "upset"],
                    ["FEED", "Temperature", 30, "C", "upset"],
                ]
            },
        ),
        "long-cases.xlsx",
    )
    streams = materials(report)
    assert {
        (s.identifier, s.case, value(s, "mass_flow").original_value) for s in streams
    } == {("FEED", "base", 10), ("FEED", "upset", 20)}
    assert_graph(report)


def test_cached_date_is_not_misrepresented_as_engineering_number():
    book = openpyxl.Workbook()
    book.active["A1"] = "=1+2"
    book.active["A1"].number_format = "yyyy-mm-dd"
    output = BytesIO()
    book.save(output)
    buffer = BytesIO()
    with (
        ZipFile(BytesIO(output.getvalue())) as src,
        ZipFile(buffer, "w", ZIP_DEFLATED) as dst,
    ):
        for item in src.infolist():
            raw = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                raw = raw.replace(b"<v />", b"<v>45000</v>").replace(
                    b"<v></v>", b"<v>45000</v>"
                )
            dst.writestr(item, raw)
    raw = buffer.getvalue()
    (FIXTURES / "date-cache.xlsx").write_bytes(raw)
    sheets, _ = read_book(raw, "date-cache.xlsx")
    cell = sheets[0].cell(1, 1)
    assert cell.formula == "=1+2" and cell.cached
    assert isinstance(cell.value, str) and cell.value.startswith("2023-")


def test_explicit_unit_ledger_connections_use_scoped_stream_keys():
    report = parse_workbook(
        workbook(
            "unit-ledger.xlsx",
            {
                "Data": [
                    ["Stream ID", "Mass flow (kg/h)"],
                    ["S-1", 100],
                    ["S-2", 90],
                    [],
                    ["Equipment tag", "Inlet streams", "Outlet streams"],
                    ["r-1", "S-1", "S-2"],
                ]
            },
        ),
        "unit-ledger.xlsx",
    )
    streams = {s.identifier: s for s in materials(report)}
    unit = next(e for e in report.entities if e.kind == "equipment")
    assert set(unit.stream_keys) == {s.key for s in streams.values()}
    assert any(
        e.stream_key == streams["S-1"].key
        and e.equipment_key == unit.key
        and e.role == "destination"
        for e in report.connections
    )
    assert any(
        e.stream_key == streams["S-2"].key
        and e.equipment_key == unit.key
        and e.role == "source"
        for e in report.connections
    )
    assert_graph(report)
