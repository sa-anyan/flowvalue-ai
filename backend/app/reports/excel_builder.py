"""Editable inputs and native Excel financial formulas; no cached financial outputs."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

from app.engine.costing import CORRELATIONS, EquipmentBOM
from app.financial.dcf_model import FinancialResult

NAVY = "162B2A"
GREEN = "167566"
CURRENCY = '"$"#,##0;("$"#,##0);"-"'
NUMBER = '#,##0.00;(#,##0.00);"-"'


def literal(cell, value):
    cell.value = str(value)
    cell.data_type = "s"  # Uploaded names must never become Excel formulas.


def build_excel(
    bom: EquipmentBOM, financial: FinancialResult, *, demo: bool = False
) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name in ("Summary", "Inputs", "Equipment BOM", "Cash Flow", "Notes"):
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_properties.tabColor = GREEN
        ws.column_dimensions["A"].width = 4
    wb.calculation = CalcProperties(
        calcId=191029, fullCalcOnLoad=True, forceFullCalc=True, calcMode="auto"
    )
    summary, inputs, equipment, cash, notes = (wb[name] for name in wb.sheetnames)
    f = financial.inputs
    for ws, title in (
        (summary, "FlowValue AI | Financial model"),
        (inputs, "Model inputs"),
        (equipment, "Equipment bill of materials"),
        (cash, "Annual cash flows (USD)"),
        (notes, "Calculation basis and scope"),
    ):
        ws["B2"] = title
        ws["B2"].font = Font(name="Arial", size=16, color=NAVY, bold=True)
        ws.row_dimensions[2].height = 30
    summary["B3"] = "Synthetic demonstration" if demo else "Current assessment scenario"
    summary["B4"] = (
        "Edit blue inputs. Financial results recalculate in Excel or LibreOffice."
    )
    summary.column_dimensions["B"].width = 32
    summary.column_dimensions["C"].width = 25
    input_values = [
        (5, "Operating life (years)", f.operating_years),
        (
            6,
            "Depreciation life (years)",
            f.depreciation_years if f.depreciation_years else "=C5",
        ),
        (7, "Discount rate", f.discount_rate),
        (8, "Tax rate", f.tax_rate),
        (9, "Plant capacity (units/year)", f.plant_capacity_units_per_year),
        (10, "Capacity factor", f.capacity_factor),
        (11, "Selling price (USD/unit)", f.product_selling_price_usd_per_unit),
        (12, "Fixed OPEX (USD/year)", f.fixed_opex_usd_per_year),
        (13, "Feedstock OPEX (USD/year)", f.feedstock_opex_usd_per_year),
        (14, "Utility OPEX (USD/year)", f.utility_opex_usd_per_year),
        (
            15,
            "Other variable OPEX (USD/year)",
            f.variable_opex_usd_per_year
            - f.feedstock_opex_usd_per_year
            - f.utility_opex_usd_per_year,
        ),
        (16, "Total variable OPEX (USD/year)", "=SUM(C13:C15)"),
        (17, "Total CAPEX (USD)", f"='Equipment BOM'!T{6 + len(bom.items)}"),
        (18, "Target CEPCI (assumption)", bom.settings.cepci_current),
        (19, "Reference CEPCI", bom.settings.cepci_base),
        (20, "Capital stress multiplier", bom.settings.capital_cost_multiplier),
    ]
    for row, label, value in input_values:
        inputs.cell(row, 2, label)
        inputs.cell(row, 3, value)
    inputs.column_dimensions["B"].width = 37
    inputs.column_dimensions["C"].width = 23
    for row in (7, 8, 10):
        inputs.cell(row, 3).number_format = "0.0%"
    inputs["B22"] = (
        "Financials: constant-dollar, end-of-year, unlevered. LCOP is pre-tax."
    )
    inputs["B23"] = "CEPCI 800/397 is configurable, not a verified current index."
    inputs["B24"] = "OPEX is annual at the selected utilization; no additional scaling."
    for cell_range, typ, low, high in (
        ("C5:C6", "whole", 1, 100),
        ("C7", "decimal", 0, 10),
        ("C8", "decimal", 0, 1),
        ("C10", "decimal", 0.000001, 1),
        ("C9", "decimal", 0.000001, 1e15),
        ("C11:C15", "decimal", 0, 1e15),
        ("C18:C19", "decimal", 0.000001, 1e12),
        ("C20", "decimal", 0.5, 1.5),
    ):
        dv = DataValidation(type=typ, operator="between", formula1=low, formula2=high)
        dv.showErrorMessage = True
        dv.errorTitle, dv.error = (
            "Invalid model input",
            f"Enter a value between {low} and {high}.",
        )
        inputs.add_data_validation(dv)
        dv.add(cell_range)
    headers = [
        "Unit",
        "Equipment",
        "Qty",
        "Material",
        "Size",
        "Size unit",
        "K1",
        "K2",
        "K3",
        "B1",
        "B2",
        "Fm",
        "Fp",
        "Base Cp0",
        "Indexed Cp0",
        "Purchased / unit",
        "Module factor",
        "CBM / unit",
        "Purchased total",
        "CBM total",
    ]
    for col, header in enumerate(headers, 1):
        equipment.cell(5, col, header)
        equipment.column_dimensions[get_column_letter(col)].width = (
            18 if col > 6 else [15, 23, 8, 22, 16, 16][col - 1]
        )
    equipment["B3"] = (
        "Edit size, quantity and factors; mechanical sizing is a snapshot. Material labels do not override Fm."
    )
    for row, item in enumerate(bom.items, 6):
        key = "vessel" if item.kind == "reactor" else item.kind
        k1, k2, k3, b1, b2, *_ = CORRELATIONS[key]
        values = [
            item.unit_id,
            item.kind,
            item.quantity,
            item.material,
            item.correlation_size,
            item.correlation_size_unit,
            k1,
            k2,
            k3,
            b1,
            b2,
            item.material_factor,
            item.pressure_factor,
        ]
        for col, value in enumerate(values, 1):
            if isinstance(value, str):
                literal(equipment.cell(row, col), value)
            else:
                equipment.cell(row, col, value)
        formulas = {
            14: f"=10^(G{row}+H{row}*LOG10(E{row})+I{row}*LOG10(E{row})^2)",
            15: f"=N{row}*Inputs!$C$18/Inputs!$C$19",
            16: f"=O{row}*L{row}*M{row}*Inputs!$C$20",
            17: f"=J{row}+K{row}*L{row}*M{row}",
            18: f"=O{row}*Q{row}*Inputs!$C$20",
            19: f"=P{row}*C{row}",
            20: f"=R{row}*C{row}",
        }
        for col, formula in formulas.items():
            equipment.cell(row, col, formula)
        for col in (14, 15, 16, 18, 19, 20):
            equipment.cell(row, col).number_format = CURRENCY
    total_row = 6 + len(bom.items)
    equipment.cell(total_row, 2, "TOTAL")
    for col in (19, 20):
        letter = get_column_letter(col)
        equipment.cell(
            total_row, col, f"=SUM({letter}6:{letter}{total_row - 1})"
        ).number_format = CURRENCY
    equipment.auto_filter.ref = f"A5:T{total_row - 1}"
    equipment.freeze_panes = "E6"
    cash_headers = [
        "Year",
        "Production",
        "Revenue",
        "Fixed OPEX",
        "Variable OPEX",
        "Total OPEX",
        "Depreciation",
        "Taxable income",
        "Taxes",
        "Net income",
        "CAPEX",
        "Free cash flow",
        "Discounted FCF",
        "Cumulative FCF",
        "Cumulative discounted FCF",
        "Discounted costs",
        "Discounted production",
        "Payback recovery year",
        "Discounted recovery year",
    ]
    for col, header in enumerate(cash_headers, 1):
        cash.cell(5, col, header)
        cash.column_dimensions[get_column_letter(col)].width = 22 if col > 1 else 9
    cash["B3"] = (
        "Years 0-100 are formula-driven. Years beyond the selected plant life contribute zero."
    )
    for row in range(6, 107):
        cash.cell(row, 1, "=0" if row == 6 else f"=A{row - 1}+1")
        if row == 6:
            for col in range(2, 20):
                cash.cell(row, col, "=0")
            for col, formula in {
                11: "=Inputs!C17",
                12: "=-K6",
                13: "=L6",
                14: "=L6",
                15: "=M6",
                16: "=K6",
            }.items():
                cash.cell(row, col, formula)
            continue
        active = f"$A{row}<=Inputs!$C$5"
        formulas = {
            2: "Inputs!$C$9*Inputs!$C$10",
            3: f"B{row}*Inputs!$C$11",
            4: "Inputs!$C$12",
            5: "Inputs!$C$16",
            6: f"SUM(D{row}:E{row})",
            7: f"IF($A{row}<=Inputs!$C$6,Inputs!$C$17/Inputs!$C$6,0)",
            8: f"C{row}-F{row}-G{row}",
            9: f"MAX(0,H{row})*Inputs!$C$8",
            10: f"H{row}-I{row}",
            11: "0",
            12: f"C{row}-F{row}-I{row}-K{row}",
            13: f"L{row}/(1+Inputs!$C$7)^$A{row}",
            16: f"(K{row}+F{row})/(1+Inputs!$C$7)^$A{row}",
            17: f"B{row}/(1+Inputs!$C$7)^$A{row}",
            18: f"IF(AND(N{row - 1}<0,N{row}>=0),A{row}-1-N{row - 1}/L{row},0)",
            19: f"IF(AND(O{row - 1}<0,O{row}>=0),A{row}-1-O{row - 1}/M{row},0)",
        }
        for col, expression in formulas.items():
            cash.cell(row, col, f"=IF({active},{expression},0)")
        cash.cell(row, 14, f"=N{row - 1}+L{row}")
        cash.cell(row, 15, f"=O{row - 1}+M{row}")
    for row in cash.iter_rows(min_row=6, max_row=106, min_col=3, max_col=16):
        for cell in row:
            cell.number_format = CURRENCY
    cash.freeze_panes = "C6"
    cash.auto_filter.ref = "A5:S106"
    metrics = [
        (
            6,
            "Net present value",
            "=NPV(Inputs!C7,'Cash Flow'!L7:L106)+'Cash Flow'!L6",
            CURRENCY,
        ),
        (
            7,
            "IRR / DCFROR",
            "=IF(MAX('Cash Flow'!L7:L106)<=0,\"n.a.\",IF(Inputs!C5=1,'Cash Flow'!L7/(-'Cash Flow'!L6)-1,IRR(OFFSET('Cash Flow'!$L$6,0,0,Inputs!$C$5+1,1),IF(SUM('Cash Flow'!L7:L106)<Inputs!C17,-0.1,0.1))))",
            "0.0%",
        ),
        (
            8,
            "Undiscounted payback (years)",
            "=IF(SUM('Cash Flow'!R7:R106)=0,\"n.a.\",SUM('Cash Flow'!R7:R106))",
            NUMBER,
        ),
        (
            9,
            "Discounted payback (years)",
            "=IF(SUM('Cash Flow'!S7:S106)=0,\"n.a.\",SUM('Cash Flow'!S7:S106))",
            NUMBER,
        ),
        (
            10,
            "LCOP (USD/unit, pre-tax)",
            "=SUM('Cash Flow'!P6:P106)/SUM('Cash Flow'!Q6:Q106)",
            '"$"0.00',
        ),
        (12, "Total CAPEX", "=Inputs!C17", CURRENCY),
        (13, "Annual OPEX", "=SUM(Inputs!C12,Inputs!C16)", CURRENCY),
        (14, "Annual production", "=Inputs!C9*Inputs!C10", "#,##0"),
    ]
    for row, label, formula, fmt in metrics:
        summary.cell(row, 2, label)
        summary.cell(row, 3, formula).number_format = fmt
    for row in range(12, 18):
        inputs.cell(row, 3).number_format = CURRENCY
    for row in range(6, 107):
        cash.cell(row, 2).number_format = "#,##0"
        cash.cell(row, 17).number_format = "#,##0"
        for col in (18, 19):
            cash.cell(row, col).number_format = NUMBER
    for row in range(6, total_row):
        equipment.cell(row, 3).number_format = "#,##0"
        equipment.cell(row, 17).number_format = "0.00"
        for col in range(7, 12):
            equipment.cell(row, col).number_format = "0.0000"
    summary["B17"] = (
        "The model contains no macros, passwords or external workbook links."
    )
    summary["B18"] = (
        "Change operating life (1-100 years) on Inputs. Cash Flow already includes all 100 years."
    )
    summary["B19"] = (
        "Risk analysis is recomputed by the app, not simulated inside this workbook."
    )
    all_notes = (
        [
            note
            for note in bom.assumptions
            if not note.startswith("Sources and validity:")
        ]
        + financial.assumptions
        + financial.warnings
    )
    for item in bom.items:
        all_notes.extend(f"{item.unit_id}: {warning}" for warning in item.warnings)
    all_notes += [
        "Cost coefficients: Turton et al., Analysis, Synthesis, and Design of Chemical Processes, Table A.1.",
        "https://github.com/pbtamarona/OpenPyTEA/blob/main/src/openpytea/data/cost_correlations.csv",
        "Module factors: Cloete et al., Journal of Cleaner Production 275 (2020), 123024.",
        "https://www.mdpi.com/1996-1073/8/4/2714",
        "Mechanical sizing, material names and pressure factors are export snapshots. Editable sizes/factors drive equipment cost formulas.",
        "All monetary values USD. No debt, working capital, tax credits, salvage or cost escalation.",
        "IRR uses the active horizon. For one operating year the equivalent return is closing cash flow / initial CAPEX - 1.",
    ]
    notes.column_dimensions["B"].width = 120
    for row, note in enumerate(all_notes, 5):
        literal(notes.cell(row, 2), note)
        notes.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="center")
        notes.row_dimensions[row].height = 32
    for ws in wb:
        for row in ws:
            for cell in row:
                if cell.value is None:
                    continue
                if cell.row == 2:
                    continue
                is_formula = cell.data_type == "f"
                color = (
                    ("00804A" if "!" in str(cell.value) else "172B2A")
                    if is_formula
                    else (
                        "2465BA" if isinstance(cell.value, (int, float)) else "364B48"
                    )
                )
                if ws == summary:
                    color = "172B2A"
                cell.font = Font(name="Arial", size=10, color=color)
                if not (ws == notes and cell.row >= 5):
                    cell.alignment = Alignment(
                        vertical="center",
                        horizontal="right"
                        if is_formula or isinstance(cell.value, (int, float))
                        else "left",
                    )
                if cell.number_format == "General" and isinstance(cell.value, float):
                    cell.number_format = NUMBER
        ws.sheet_properties.outlinePr.summaryRight = False
    for ws, end in ((equipment, 20), (cash, 19)):
        ws.row_dimensions[5].height = 34
        for cell in ws[5][:end]:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        for row in ws.iter_rows(min_row=6):
            for cell in row[:end]:
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F1F6F4")
        ws.print_options.horizontalCentered = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:5"
    for ws, row, cols in (
        (equipment, total_row, range(1, 21)),
        (summary, 6, range(2, 4)),
    ):
        for col in cols:
            ws.cell(row, col).border = Border(top=Side(style="thin", color=GREEN))
            ws.cell(row, col).font = Font(name="Arial", size=11, bold=True, color=NAVY)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
