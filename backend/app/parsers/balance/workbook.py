"""Bounded workbook reading; never evaluates formula text or follows links."""

import posixpath
import re
from dataclasses import dataclass, field
from io import BytesIO, StringIO
from pathlib import PurePath
from zipfile import BadZipFile, ZipFile

import openpyxl
from defusedxml.ElementTree import fromstring
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import coordinate_to_tuple, get_column_letter, range_boundaries
from openpyxl.utils.datetime import from_excel

MAX_FILE = 4 * 1024 * 1024
MAX_CELLS = 150_000
MAX_ROWS = 10_000
MAX_COLS = 256
MAX_SHEETS = 40


@dataclass
class Cell:
    value: object = None
    formula: str | None = None
    error: bool = False
    cached: bool = False
    number_format: str = ""


@dataclass
class Sheet:
    name: str
    cells: dict[tuple[int, int], Cell]
    rows: int
    cols: int
    hidden: bool = False
    merged: dict[tuple[int, int], tuple[int, int]] = field(default_factory=dict)

    def cell(self, r, c):
        return self.cells.get((r, c), Cell())

    def text(self, r, c):
        v = self.cell(r, c).value
        return str(v if v is not None else "").strip()

    def header(self, r, c):
        rr, cc = self.merged.get((r, c), (r, c))
        return self.cell(rr, cc), (rr, cc)


def address(r, c):
    return f"{get_column_letter(c)}{r}"


def safe_filename(name):
    return (
        re.sub(r"[\x00-\x1f]", "", str(name).replace("\\", "/").split("/")[-1])[:255]
        or "workbook.xlsx"
    )


def read_book(content: bytes, name: str):
    if not content or len(content) > MAX_FILE:
        raise ValueError("Workbook must be non-empty and at most 4 MiB.")
    ext = PurePath(name.lower()).suffix
    warnings = []
    if ext in (".xlsx", ".xlsm"):
        if not content.startswith(b"PK"):
            raise ValueError("The extension does not match an OOXML workbook.")
        try:
            with ZipFile(BytesIO(content)) as z:
                entries = z.infolist()
                if (
                    len(entries) > 2048
                    or sum(i.file_size for i in entries) > 40 * 1024 * 1024
                    or any(
                        i.flag_bits & 1
                        or i.file_size > 20 * 1024 * 1024
                        or i.file_size > max(i.compress_size, 1) * 1000
                        for i in entries
                    )
                ):
                    raise ValueError(
                        "Workbook archive exceeds safe expansion limits or is encrypted."
                    )
                if (
                    "[Content_Types].xml" not in z.namelist()
                    or "xl/workbook.xml" not in z.namelist()
                ):
                    raise ValueError("Not an Excel OOXML workbook.")
                if any("vbaproject" in i.filename.lower() for i in entries):
                    if ext != ".xlsm":
                        raise ValueError("Macro-bearing workbooks require the .xlsm format.")
                    warnings.append("VBA macros were not executed or imported; only stored spreadsheet values were read.")
                if len({i.filename for i in entries}) != len(entries):
                    raise ValueError("Duplicate archive members are not supported.")
                actual_cells = 0
                caches = {}
                merged_ranges = {}
                for i in entries:
                    if i.filename.endswith((".xml", ".rels")):
                        xml = z.read(i)
                        if b"<!DOCTYPE" in xml.upper() or b"<!ENTITY" in xml.upper():
                            raise ValueError(
                                "DTD/entity declarations are not supported."
                            )
                        if "externalLinks/" in i.filename:
                            warnings.append(
                                "External workbook links were not fetched; dependent caches are unverified."
                            )
                        root = fromstring(
                            xml
                        )  # reject malformed/unsafe XML before the OOXML library reads it
                        if i.filename.startswith(
                            "xl/worksheets/"
                        ) and i.filename.endswith(".xml"):
                            for node in root.iter():
                                tag = node.tag.rsplit("}", 1)[-1]
                                if tag == "c":
                                    actual_cells += 1
                                    rr, cc = coordinate_to_tuple(
                                        node.attrib.get("r", "A1")
                                    )
                                    if any(child.tag.rsplit("}", 1)[-1] == "f" for child in node):
                                        value = next((child.text for child in node if child.tag.rsplit("}", 1)[-1] == "v"), None)
                                        kind = node.attrib.get("t", "n")
                                        if value is not None:
                                            if kind == "b":
                                                if value not in ("0", "1"):
                                                    raise ValueError("Malformed boolean formula cache; workbook not parsed.")
                                                value = value == "1"
                                            elif kind not in ("str", "e", "s"):
                                                try:
                                                    value = float(value) if any(ch in value.lower() for ch in (".", "e")) else int(value)
                                                except ValueError as exc:
                                                    raise ValueError("Malformed numeric formula cache; workbook not parsed.") from exc
                                        if kind == "s":
                                            value = None  # unsupported shared-string cache index is not a literal value
                                            warnings.append("A shared-string formula cache could not be resolved; formula evidence retained.")
                                        caches.setdefault(i.filename, {})[rr, cc] = (value, kind == "e")
                                    if (
                                        rr > MAX_ROWS
                                        or cc > MAX_COLS
                                        or actual_cells > MAX_CELLS
                                    ):
                                        raise ValueError(
                                            "Workbook exceeds safe worksheet limits."
                                        )
                                if tag in (
                                    "dimension",
                                    "mergeCell",
                                ) and node.attrib.get("ref"):
                                    if tag == "mergeCell":
                                        merged_ranges.setdefault(i.filename, []).append(node.attrib["ref"])
                                    c0, r0, c1, r1 = range_boundaries(
                                        node.attrib["ref"]
                                    )
                                    if (
                                        r1 > MAX_ROWS
                                        or c1 > MAX_COLS
                                        or (r1 - r0 + 1) * (c1 - c0 + 1) > MAX_CELLS
                                    ):
                                        raise ValueError(
                                            "Worksheet dimensions/merged range exceed safe limits."
                                        )
                        if (
                            i.filename == "xl/styles.xml"
                            and sum(1 for _ in root.iter()) > 20000
                        ):
                            raise ValueError(
                                "Workbook style definitions exceed safe limits."
                            )
                workbook_xml = fromstring(z.read("xl/workbook.xml"))
                relations = fromstring(z.read("xl/_rels/workbook.xml.rels"))
                targets = {item.attrib["Id"]: item.attrib.get("Target", "") for item in relations
                           if item.attrib.get("TargetMode") != "External"}
                worksheet_paths = {}
                for node in workbook_xml.iter():
                    if node.tag.rsplit("}", 1)[-1] == "sheet":
                        rel = next((v for k, v in node.attrib.items() if k.endswith("}id")), None)
                        target = targets.get(rel, "")
                        worksheet_paths[node.attrib["name"]] = (target.lstrip("/") if target.startswith("/")
                            else posixpath.normpath(posixpath.join("xl", target)))
            # The validated OOXML scan already captures formula caches and merges.
            # One read-only library pass retains formula text (including shared
            # formula expansion) and styles. No second workbook open is needed.
            formulas = openpyxl.load_workbook(
                BytesIO(content), data_only=False, read_only=True, keep_links=False
            )
            sheets = []
            total = 0
            try:
                if len(formulas.worksheets) > MAX_SHEETS:
                    raise ValueError("Workbook exceeds 40 worksheets.")
                for ws in formulas:
                    rows, cols = ws.max_row, ws.max_column
                    total += rows * cols
                    if rows > MAX_ROWS or cols > MAX_COLS or total > MAX_CELLS:
                        raise ValueError(
                            "Workbook exceeds 10,000 rows, 256 columns or 150,000 scanned cells."
                        )
                    cells = {}
                    merges = {}
                    path = worksheet_paths.get(ws.title, "")
                    formula_caches = caches.get(path, {})
                    for row in ws.iter_rows():
                        for cell in row:
                            formula = str(cell.value) if cell.data_type == "f" else None
                            if formula and len(formula) > 2000:
                                raise ValueError("Formula exceeds 2,000 characters.")
                            if cell.value is None and not formula:
                                continue
                            value, error = formula_caches.get((cell.row, cell.column), (None, False)) if formula else (cell.value, cell.data_type == "e")
                            if formula and isinstance(value, (int, float)) and not isinstance(value, bool) and is_date_format(cell.number_format):
                                value = from_excel(value, formulas.epoch)
                            if not isinstance(value, (int, float, str, bool, type(None))):
                                value = str(value)
                            if isinstance(value, str) and len(value) > 4000:
                                raise ValueError("Cell text exceeds 4,000 characters.")
                            cells[cell.row, cell.column] = Cell(value, formula, error, bool(formula), cell.number_format)
                    for ref in merged_ranges.get(path, []):
                        c0, r0, c1, r1 = range_boundaries(ref)
                        for rr in range(r0, r1 + 1):
                            for cc in range(c0, c1 + 1):
                                merges[rr, cc] = (r0, c0)
                    sheets.append(
                        Sheet(
                            ws.title,
                            cells,
                            rows,
                            cols,
                            ws.sheet_state != "visible",
                            merges,
                        )
                    )
            finally:
                formulas.close()
            return sheets, list(dict.fromkeys(warnings))
        except (BadZipFile, KeyError, TypeError, OSError, IndexError) as e:
            raise ValueError("Malformed or unsupported Excel workbook.") from e
    if ext == ".xls":
        import xlrd

        if not content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            raise ValueError(
                "Only OLE/BIFF .xls files are supported; renamed HTML/XML files are not."
            )
        try:
            book = xlrd.open_workbook(
                file_contents=content,
                formatting_info=True,
                on_demand=True,
                ragged_rows=True,
                logfile=StringIO(),
            )
            sheets = []
            total = 0
            try:
                if book.nsheets > MAX_SHEETS:
                    raise ValueError("Workbook exceeds 40 worksheets.")
                for i in range(book.nsheets):
                    ws = book.sheet_by_index(i)
                    total += ws.nrows * ws.ncols
                    if ws.nrows > MAX_ROWS or ws.ncols > MAX_COLS or total > MAX_CELLS:
                        raise ValueError("Workbook exceeds safe cell limits.")
                    cells = {}
                    merges = {}
                    for r in range(ws.nrows):
                        for c in range(ws.row_len(r)):
                            v = ws.cell(r, c)
                            if v.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                                continue
                            value = (
                                xlrd.error_text_from_code.get(v.value, "#ERROR!")
                                if v.ctype == xlrd.XL_CELL_ERROR
                                else v.value
                            )
                            if isinstance(value, str) and len(value) > 4000:
                                raise ValueError("Cell text exceeds 4,000 characters.")
                            fmt = book.format_map[
                                book.xf_list[ws.cell_xf_index(r, c)].format_key
                            ].format_str
                            cells[r + 1, c + 1] = Cell(
                                value, None, v.ctype == xlrd.XL_CELL_ERROR, True, fmt
                            )
                    for r0, r1, c0, c1 in ws.merged_cells:
                        for r in range(r0, r1):
                            for c in range(c0, c1):
                                merges[r + 1, c + 1] = (r0 + 1, c0 + 1)
                    sheets.append(
                        Sheet(
                            ws.name,
                            cells,
                            ws.nrows,
                            ws.ncols,
                            bool(ws.visibility),
                            merges,
                        )
                    )
                    book.unload_sheet(i)
            finally:
                book.release_resources()
            return sheets, [
                "Legacy .xls exposes stored results only; formula text/cache age cannot be determined. Extracted values require confirmation."
            ]
        except xlrd.XLRDError as e:
            raise ValueError(
                "Malformed, encrypted or unsupported BIFF workbook."
            ) from e
    raise ValueError(
        "Choose an .xlsx, .xls or .xlsm workbook. Password-protected files, .xlsb and renamed formats are not supported."
    )
